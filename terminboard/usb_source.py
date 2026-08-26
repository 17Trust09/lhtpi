"""Terminboard – USB-Stick mit ``termine.xlsx`` (oder ``termine.csv``) als Terminquelle.

Konvention
----------
Der Stick-Root enthält eine Datei ``termine.xlsx`` (bevorzugt) oder ``termine.csv``
(Fallback). Beide liefern Termine mit denselben sechs Feldern.

Excel-Vorlage (``USB/termine.xlsx``), Spalten::

    Art | Titel | Referenz | Von | Bis | Hinweis

* ``Art``      – Kalibrierung | Audit | Wartung | Info (sonst → info)
* ``Titel``    – Pflicht (Zeilen ohne Titel werden ignoriert)
* ``Referenz`` – optional (z. B. ``P3``)
* ``Von``      – Startdatum ``TT.MM.JJJJ`` oder ``JJJJ-MM-TT`` (Pflicht)
* ``Bis``      – optional, gleiche Datumsformate
* ``Hinweis``  – optionaler Einzeiler

CSV-Variante (``termine.csv``), Semikolon-getrennt, Header-Zeile::

    typ;titel;referenz;start;ende;text

Leer- und Kommentarzeilen (``#``) werden ignoriert. Zeilen ohne gültiges
Startdatum, ohne Titel oder mit zu wenigen Spalten werden übersprungen.
"""
import csv
import os
from datetime import date, datetime

from models import db, Setting

ALLOWED_TYPES = {'kalibrierung', 'audit', 'info', 'wartung'}

EXPECTED_HEADER = ['typ', 'titel', 'referenz', 'start', 'ende', 'text']

USB_CSV_FILENAME = 'termine.csv'
USB_XLSX_FILENAME = 'termine.xlsx'
USB_TERMIN_FILENAMES = (USB_XLSX_FILENAME, USB_CSV_FILENAME)  # xlsx hat Vorrang
FIXED_MOUNT = '/mnt/lhtpi-usb'  # gemeinsamer Mount-Point mit LHTPi (ein Stick für beide Apps)

# Verständliche Spaltenüberschriften → interner Schlüssel (für die Kopfzeilen-Erkennung)
HEADER_ALIASES = {
    'art': 'typ', 'typ': 'typ', 'type': 'typ',
    'titel': 'titel', 'title': 'titel',
    'referenz': 'referenz', 'reference': 'referenz', 'ref': 'referenz',
    'von': 'start', 'start': 'start',
    'bis': 'ende', 'ende': 'ende', 'end': 'ende',
    'hinweis': 'text', 'text': 'text', 'bemerkung': 'text', 'notiz': 'text',
}

# Einstellungs-Keys
SETTING_MODE = 'player_mode'            # 'auto' | 'manual'
SETTING_MANUAL_SOURCE = 'manual_source'  # 'web' | 'usb'


# ── Reine Helfer (ohne DB, gut testbar) ───────────────────────────────

def parse_date(value):
    """Parst ``TT.MM.JJJJ`` oder ``JJJJ-MM-TT``. Liefert ``date`` oder ``None``."""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def normalize_type(value):
    """Normalisiert die Art auf ``kalibrierung|audit|wartung|info`` (sonst ``info``)."""
    t = str(value or '').strip().lower()
    return t if t in ALLOWED_TYPES else 'info'


def _cell_to_date(value):
    """Wandelt einen Excel-Zellenwert (``date``/``datetime``/String) in ``date``."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(value)


def _looks_like_header(cells):
    """True, wenn die erste Zelle wie eine Spaltenüberschrift aussieht."""
    if not cells:
        return False
    return str(cells[0] or '').strip().lower() in HEADER_ALIASES


def parse_termin_csv(text):
    """Parst CSV-Text in eine Liste von Termin-Dicts (``date``-Objekte).

    Verwendet das ``csv``-Modul (unterstützt Quoting), erkennt die Header-Zeile
    und überspringt Zeilen ohne gültiges Pflicht-Startdatum oder ohne Titel.
    """
    rows = []
    lines = [l for l in (x.strip() for x in text.splitlines())
             if l and not l.startswith('#')]
    if not lines:
        return rows

    reader = csv.reader(lines, delimiter=';')
    first = next(reader, None)
    if first is None:
        return rows

    is_header = [c.strip().lower() for c in first] == EXPECTED_HEADER
    records = list(reader)
    if not is_header:
        # Keine Header-Zeile vorhanden → erste Zeile ist ein Datensatz
        records.insert(0, first)

    for parts in records:
        if len(parts) < 4:
            continue  # zu wenige Spalten (mind. typ;titel;referenz;start)
        while len(parts) < 6:
            parts.append('')
        typ = normalize_type(parts[0])
        titel = parts[1].strip()
        referenz = parts[2].strip() or None
        start = parse_date(parts[3].strip())
        ende = parse_date(parts[4].strip())
        text = parts[5].strip() or None
        if not titel:
            continue
        if start is None:
            continue  # Pflicht-Startdatum fehlt oder ist ungültig
        rows.append({
            'typ': typ,
            'titel': titel,
            'referenz': referenz,
            'start': start,
            'ende': ende,
            'text': text,
        })
    return rows


def parse_termin_xlsx(file_path):
    """Liest ``termine.xlsx`` (openpyxl) in Termin-Dicts.

    Spalten positionell: Art, Titel, Referenz, Von, Bis, Hinweis. Die erste
    Datenzeile wird als Kopfzeile erkannt und übersprungen. Datumszellen
    werden als ``date`` übernommen.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    try:
        wb = load_workbook(file_path, data_only=True, read_only=True)
    except Exception:
        return []

    rows = []
    try:
        ws = wb['Termine'] if 'Termine' in wb.sheetnames else wb.worksheets[0]
        header_seen = False
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == '' for c in row):
                continue
            if not header_seen and _looks_like_header(row):
                header_seen = True
                continue
            header_seen = True
            parts = list(row[:6])
            while len(parts) < 6:
                parts.append('')
            typ = normalize_type(parts[0])
            titel = str(parts[1] or '').strip()
            referenz = str(parts[2] or '').strip() or None
            start = _cell_to_date(parts[3])
            ende = _cell_to_date(parts[4])
            text = str(parts[5] or '').strip() or None
            if not titel:
                continue
            if start is None:
                continue
            rows.append({
                'typ': typ,
                'titel': titel,
                'referenz': referenz,
                'start': start,
                'ende': ende,
                'text': text,
            })
    finally:
        wb.close()
    return rows


def read_csv_from_dir(mount_dir):
    """Liest ``termine.csv`` aus ``mount_dir`` und liefert Termin-Dicts."""
    path = os.path.join(mount_dir, USB_CSV_FILENAME)
    if not os.path.isfile(path):
        return []
    try:
        with open(path, encoding='utf-8', errors='ignore') as fh:
            return parse_termin_csv(fh.read())
    except OSError:
        return []


def read_termin_from_dir(mount_dir):
    """Liest Termine aus ``mount_dir`` — bevorzugt ``termine.xlsx``, sonst ``termine.csv``."""
    xlsx = os.path.join(mount_dir, USB_XLSX_FILENAME)
    if os.path.isfile(xlsx):
        return parse_termin_xlsx(xlsx)
    return read_csv_from_dir(mount_dir)


# ── Einstellungs-Helfer (DB) ──────────────────────────────────────────

def get_setting(key, default=None):
    row = db.session.get(Setting, key)
    return row.value if row else default


def set_setting(key, value):
    row = db.session.get(Setting, key)
    if row is None:
        row = Setting(key=key, value='')
        db.session.add(row)
    row.value = str(value)
    db.session.commit()


def get_mode():
    return get_setting(SETTING_MODE, 'auto')


def get_manual_source():
    return get_setting(SETTING_MANUAL_SOURCE, 'web')


# ── USB-Erkennung ─────────────────────────────────────────────────────

def _is_valid_dir(path):
    """True, wenn ``path`` ein lesbarer Ordner ist (schützt vor stale Mount)."""
    if not os.path.isdir(path):
        return False
    try:
        os.listdir(path)
        return True
    except OSError:
        return False


def _iter_mount_candidates():
    """Mögliche Pfade, unter denen ein Stick (mit ``termine.xlsx``/``termine.csv``) liegen kann.

    Unter ``/media`` und ``/run/media`` wird zusätzlich eine Ebene tiefer
    gesucht (Auto-Mount-Layout ``/media/<user>/<label>``). Unter ``/mnt``
    liegen Mount-Points direkt als Kinder (z. B. ``/mnt/lhtpi-usb``) — hier
    wird nicht tiefer gesucht, damit die Termin-Datei nicht fälschlich aus
    Unterordnern wie ``slides/`` erkannt wird.
    """
    seen = {FIXED_MOUNT}
    yield FIXED_MOUNT
    for root, deep in (('/media', True), ('/run/media', True), ('/mnt', False)):
        if not os.path.isdir(root):
            continue
        try:
            entries = os.listdir(root)
        except OSError:
            continue
        for name in entries:
            userdir = os.path.join(root, name)
            if not os.path.isdir(userdir):
                continue
            if userdir not in seen:
                yield userdir
                seen.add(userdir)
            if not deep:
                continue
            try:
                subs = os.listdir(userdir)
            except OSError:
                continue
            for sub in subs:
                cand = os.path.join(userdir, sub)
                if not os.path.isdir(cand):
                    continue
                if cand not in seen:
                    yield cand
                    seen.add(cand)


def find_usb_termin_dir():
    """Liefert den Ordner mit ``termine.xlsx`` oder ``termine.csv`` (oder ``None``)."""
    for cand in _iter_mount_candidates():
        if not _is_valid_dir(cand):
            continue
        for fname in USB_TERMIN_FILENAMES:
            if os.path.isfile(os.path.join(cand, fname)):
                return cand
    return None


def usb_termine(mount_dir=None):
    """Liefert die Termin-Dicts vom Stick (oder ``[]`` wenn keiner erkannt)."""
    mount_dir = mount_dir or find_usb_termin_dir()
    if not mount_dir:
        return []
    return read_termin_from_dir(mount_dir)
