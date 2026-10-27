"""Abgeschlossene Termine ins Blatt ``Archiv`` verschieben.

Warum das Geraet das selbst macht: Wer jeden Monat von Hand Zeilen verschiebt,
vergisst es irgendwann - und dann stehen erledigte Kalibrierungen weiter als
offen im Kiosk.

Regeln:

* Eine Zeile wandert nur, wenn der Status ``abgeschlossen`` ist.
* Die Nummer bleibt - der Pi erkennt Termine an der Nummer, nicht an der Zeile.
* Die Zeile wird aus ``Termine`` geloescht, es bleiben keine Luecken.
* Vor dem Schreiben entsteht eine Sicherung ``<datei>.bak``.
* Schlaegt etwas fehl, bleibt die Datei unveraendert. Ein halb geschriebenes
  Excel waere schlimmer als ein nicht verschobener Eintrag.
"""
import os
import shutil

TERMINE_BLATT = 'Termine'
ARCHIV_BLATT = 'Archiv'
SPALTE_STATUS = 7            # G
SPALTE_NR = 1                # A


def _hat_archiv(wb):
    return ARCHIV_BLATT in wb.sheetnames


def verschiebe_abgeschlossene(pfad, trocken=False):
    """Verschiebt abgeschlossene Zeilen ins Archiv.

    Liefert ``{'verschoben': [Nr, ...], 'fehler': str|None, 'sicherung': str|None}``.
    """
    ergebnis = {'verschoben': [], 'fehler': None, 'sicherung': None}

    if not os.path.isfile(pfad):
        ergebnis['fehler'] = 'Datei nicht gefunden'
        return ergebnis

    try:
        from openpyxl import load_workbook
        wb = load_workbook(pfad)
    except Exception as fehler:                      # noqa: BLE001
        ergebnis['fehler'] = 'Datei nicht lesbar: %s' % fehler
        return ergebnis

    if TERMINE_BLATT not in wb.sheetnames:
        ergebnis['fehler'] = 'Blatt "Termine" fehlt'
        return ergebnis

    ws = wb[TERMINE_BLATT]
    if not _hat_archiv(wb):
        archiv = wb.create_sheet(ARCHIV_BLATT)
        archiv.append([ws.cell(row=1, column=spalte).value
                       for spalte in range(1, ws.max_column + 1)])
    else:
        archiv = wb[ARCHIV_BLATT]

    # Welche Nummern stehen schon im Archiv? Doppelt archivieren waere
    # verwirrend, auch wenn die Historie im Geraet davon unberuehrt bleibt.
    schon_archiviert = set()
    for zeile in archiv.iter_rows(min_row=2, max_col=SPALTE_NR, values_only=True):
        if zeile and zeile[0] is not None:
            schon_archiviert.add(zeile[0])

    # Von unten nach oben: Loeschen verschiebt sonst die Indizes darunter.
    zu_loeschen = []
    for zeile in range(2, ws.max_row + 1):
        status = ws.cell(row=zeile, column=SPALTE_STATUS).value
        nr = ws.cell(row=zeile, column=SPALTE_NR).value
        if str(status or '').strip().lower() != 'abgeschlossen':
            continue
        if nr is None or nr in schon_archiviert:
            continue
        zu_loeschen.append(zeile)

    if not zu_loeschen:
        return ergebnis

    # Erst die Werte ins Archiv schreiben, dann loeschen - andersherum waeren
    # sie bei einem Fehler weg.
    for zeile in zu_loeschen:
        werte = [ws.cell(row=zeile, column=spalte).value
                 for spalte in range(1, ws.max_column + 1)]
        archiv.append(werte)
        ergebnis['verschoben'].append(werte[SPALTE_NR - 1])

    if trocken:
        return ergebnis

    for zeile in reversed(zu_loeschen):
        ws.delete_rows(zeile, 1)

    sicherung = pfad + '.bak'
    try:
        shutil.copy2(pfad, sicherung)
        ergebnis['sicherung'] = sicherung
    except OSError as fehler:
        ergebnis['fehler'] = 'Sicherung nicht moeglich: %s' % fehler
        return ergebnis

    try:
        wb.save(pfad)
    except Exception as fehler:                      # noqa: BLE001
        # Zurueckrollen: die Sicherung ist die Wahrheit.
        try:
            shutil.copy2(sicherung, pfad)
            ergebnis['fehler'] = 'Schreiben fehlgeschlagen, zurueckgerollt: %s' % fehler
        except OSError:
            ergebnis['fehler'] = 'Schreiben fehlgeschlagen: %s' % fehler
        ergebnis['verschoben'] = []
        return ergebnis

    return ergebnis
