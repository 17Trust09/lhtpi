"""Tests für den Excel-Parser (Modell 2)."""
import os
from datetime import datetime

import openpyxl
import pytest

from excel_modell2 import lies_datei, _zu_datetime

VORLAGE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       'vorlagen', 'termine_modell2.xlsx')


class TestDatum:
    def test_datum_ohne_zeit_wird_tagesbeginn(self):
        assert _zu_datetime('16.11.2026') == datetime(2026, 11, 16, 0, 0)

    def test_datum_als_ende_wird_tagesende(self):
        assert _zu_datetime('16.11.2026', als_ende=True) == datetime(2026, 11, 16, 23, 59)

    def test_datum_mit_zeit(self):
        assert _zu_datetime('16.11.2026 08:30') == datetime(2026, 11, 16, 8, 30)

    def test_iso_format(self):
        assert _zu_datetime('2026-11-16 08:30') == datetime(2026, 11, 16, 8, 30)

    def test_leer_ergibt_none(self):
        assert _zu_datetime('') is None
        assert _zu_datetime(None) is None

    def test_unsinn_ergibt_none(self):
        assert _zu_datetime('keine Ahnung') is None


@pytest.mark.skipif(not os.path.exists(VORLAGE), reason='Vorlage fehlt')
class TestVorlage:
    @pytest.fixture(scope='class')
    def daten(self):
        return lies_datei(VORLAGE)

    def test_termine_der_vorlage(self, daten):
        # Die Vorlage gehoert dem Betrieb - geprueft wird, dass jede Zeile
        # gelesen wird, nicht eine feste Anzahl.
        assert len(daten['zeilen']) >= 1
        assert daten['fehler'] == []

    def test_erste_zeile(self, daten):
        z = daten['zeilen'][0]
        assert z['nr'] == 1
        assert z['art'] in ('Kalibrierung', 'Wartung', 'Ausfall', 'Audit', 'Info')
        assert z['pruefstand']
        # Nur Datum im Blatt „Termine“: Von ist Tagesbeginn, Bis Tagesende.
        assert z['von'].hour == 0 and z['von'].minute == 0
        assert z['bis'].hour == 23 and z['bis'].minute == 59
        assert z['status'] in ('offen', 'abgeschlossen', 'storniert')

    def test_bloecke_der_vorlage_sind_leer(self, daten):
        # Das Blatt „Zeiten“ ist bewusst leer - Beispielblöcke würden als
        # echte Ausfallzeiten importiert.
        assert daten['bloecke'] == {}

    def test_bloecke_werden_gelesen(self, tmp_path):
        # Die Blocklogik wird an einer eigenen Datei geprüft, nicht an der
        # Vorlage: die bleibt leer, damit nichts Falsches importiert wird.
        pfad = tmp_path / 'mit_bloecken.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Termine'
        ws.append(['Nr', 'Art', 'Prüfstand', 'Von', 'Bis', 'Info', 'Status'])
        ws.append([2, 'Kalibrierung', 'P10 PTU TESTFUCHS', '16.11.2026',
                   '21.11.2026', 'unterbrochen', 'offen'])
        zeiten = wb.create_sheet('Zeiten')
        zeiten.append(['Termin (Nr)', 'Von', 'Bis', 'Notiz'])
        zeiten.append([2, datetime(2026, 11, 16, 8, 0), datetime(2026, 11, 16, 12, 0), 'Tag 1'])
        zeiten.append([2, datetime(2026, 11, 17, 8, 0), datetime(2026, 11, 17, 16, 0), 'Tag 2'])
        wb.save(pfad)

        daten = lies_datei(str(pfad))
        assert 2 in daten['bloecke']
        assert len(daten['bloecke'][2]) == 2
        assert daten['bloecke'][2][0] == (datetime(2026, 11, 16, 8, 0),
                                          datetime(2026, 11, 16, 12, 0))

    def test_pruefstaende_gelesen(self, daten):
        # Die Stammliste waechst mit dem Betrieb - geprueft wird der Inhalt,
        # nicht die Anzahl.
        assert 'P01 NRA' in daten['pruefstaende']
        assert 'P43 LOW-PRESSURE' in daten['pruefstaende']
        assert len(daten['pruefstaende']) >= 15
        assert 'P03 NRA' in daten['pruefstaende']
        assert 'P43 LOW-PRESSURE' in daten['pruefstaende']

    def test_keine_fehler(self, daten):
        assert daten['fehler'] == []


class TestFehlerzeilen:
    def _datei(self, tmp_path, termin_zeilen, zeiten_zeilen=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Termine'
        ws.append(['Nr', 'Art', 'Prüfstand', 'Von', 'Bis', 'Info', 'Status'])
        for z in termin_zeilen:
            ws.append(z)
        if zeiten_zeilen is not None:
            wz = wb.create_sheet('Zeiten')
            wz.append(['Termin', 'Von', 'Bis', 'Notiz'])
            for z in zeiten_zeilen:
                wz.append(z)
        pfad = tmp_path / 'test.xlsx'
        wb.save(pfad)
        return pfad

    def test_fehlende_nr_nimmt_zeilennummer(self, tmp_path):
        pfad = self._datei(tmp_path, [
            [None, 'Wartung', 'P03 NRA', '02.11.2026 07:00', '02.11.2026 15:00', 'x', 'offen'],
        ])
        daten = lies_datei(str(pfad))
        assert daten['zeilen'][0]['nr'] == 2   # Kopfzeile ist Zeile 1

    def test_leerer_status_wird_offen(self, tmp_path):
        pfad = self._datei(tmp_path, [
            [1, 'Wartung', 'P03 NRA', '02.11.2026 07:00', '02.11.2026 15:00', 'x', None],
        ])
        daten = lies_datei(str(pfad))
        assert daten['zeilen'][0]['status'] == 'offen'

    def test_zeiten_mit_bis_vor_von_werden_gemeldet(self, tmp_path):
        pfad = self._datei(
            tmp_path,
            [[1, 'Kalibrierung', 'P03 NRA', '16.11.2026 08:00', '16.11.2026 16:00', 'x', 'offen']],
            [[1, '16.11.2026 12:00', '16.11.2026 08:00', '']],
        )
        daten = lies_datei(str(pfad))
        assert len(daten['fehler']) == 1
        assert daten['bloecke'] == {}

    def test_zeiten_ohne_nr_gemeldet(self, tmp_path):
        pfad = self._datei(
            tmp_path,
            [[1, 'Kalibrierung', 'P03 NRA', '16.11.2026 08:00', '16.11.2026 16:00', 'x', 'offen']],
            [[None, '16.11.2026 08:00', '16.11.2026 12:00', '']],
        )
        daten = lies_datei(str(pfad))
        assert len(daten['fehler']) == 1
