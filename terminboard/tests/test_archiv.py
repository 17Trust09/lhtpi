"""Tests für das Verschieben abgeschlossener Termine ins Archiv."""
import os

import pytest
from openpyxl import Workbook, load_workbook

from archiv import verschiebe_abgeschlossene

KOPF = ['Nr', 'Art', 'Prüfstand', 'Von', 'Bis', 'Info', 'Status']


def _datei(tmp_path, zeilen, mit_archiv=True):
    pfad = tmp_path / 'termine.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Termine'
    ws.append(KOPF)
    for zeile in zeilen:
        ws.append(zeile)
    if mit_archiv:
        archiv = wb.create_sheet('Archiv')
        archiv.append(KOPF)
    wb.save(pfad)
    return str(pfad)


def test_abgeschlossene_wandern_ins_archiv(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Kalibrierung', 'P11', '19.10.2026', '26.10.2026', 'x', 'offen'],
        [2, 'Wartung', 'P03', '02.11.2026', '02.11.2026', 'y', 'abgeschlossen'],
    ])
    ergebnis = verschiebe_abgeschlossene(pfad)
    assert ergebnis['verschoben'] == [2]
    assert ergebnis['fehler'] is None

    wb = load_workbook(pfad)
    termine = [r[0] for r in wb['Termine'].iter_rows(min_row=2, values_only=True)]
    archiv = [r[0] for r in wb['Archiv'].iter_rows(min_row=2, values_only=True)]
    assert termine == [1]
    assert archiv == [2]


def test_offene_bleiben_stehen(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Kalibrierung', 'P11', '19.10.2026', '26.10.2026', 'x', 'offen'],
        [2, 'Kalibrierung', 'P43', '05.11.2026', None, 'y', 'offen'],
    ])
    ergebnis = verschiebe_abgeschlossene(pfad)
    assert ergebnis['verschoben'] == []
    wb = load_workbook(pfad)
    assert wb['Termine'].max_row == 3


def test_zeilen_werden_wirklich_geloescht(tmp_path):
    # Keine Leerzeilen: nach dem Verschieben ruecken die Zeilen zusammen.
    pfad = _datei(tmp_path, [
        [1, 'Wartung', 'P01', '02.11.2026', '02.11.2026', 'a', 'abgeschlossen'],
        [2, 'Kalibrierung', 'P11', '19.10.2026', '26.10.2026', 'b', 'offen'],
        [3, 'Wartung', 'P02', '03.11.2026', '03.11.2026', 'c', 'abgeschlossen'],
        [4, 'Kalibrierung', 'P43', '05.11.2026', None, 'd', 'offen'],
    ])
    verschiebe_abgeschlossene(pfad)
    wb = load_workbook(pfad)
    werte = [r for r in wb['Termine'].iter_rows(min_row=2, values_only=True)]
    assert [r[0] for r in werte] == [2, 4]
    assert all(r[0] is not None for r in werte)


def test_zweiter_lauf_aendert_nichts(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Wartung', 'P01', '02.11.2026', '02.11.2026', 'a', 'abgeschlossen'],
    ])
    verschiebe_abgeschlossene(pfad)
    zweiter = verschiebe_abgeschlossene(pfad)
    assert zweiter['verschoben'] == []


def test_nummer_bleibt_erhalten(tmp_path):
    # Die Nummer verbindet Termin und Zeitbloecke - sie darf sich nicht aendern.
    pfad = _datei(tmp_path, [
        [7, 'Ausfall', 'P21', '18.11.2026', '19.11.2026', 'w', 'abgeschlossen'],
    ])
    verschiebe_abgeschlossene(pfad)
    wb = load_workbook(pfad)
    assert wb['Archiv'].cell(row=2, column=1).value == 7


def test_sicherung_wird_angelegt(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Wartung', 'P01', '02.11.2026', '02.11.2026', 'a', 'abgeschlossen'],
    ])
    ergebnis = verschiebe_abgeschlossene(pfad)
    assert ergebnis['sicherung']
    assert os.path.exists(ergebnis['sicherung'])


def test_archivblatt_wird_angelegt(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Wartung', 'P01', '02.11.2026', '02.11.2026', 'a', 'abgeschlossen'],
    ], mit_archiv=False)
    verschiebe_abgeschlossene(pfad)
    wb = load_workbook(pfad)
    assert 'Archiv' in wb.sheetnames
    assert wb['Archiv'].cell(row=2, column=1).value == 1


def test_fehlende_datei_meldet_fehler(tmp_path):
    ergebnis = verschiebe_abgeschlossene(str(tmp_path / 'gibtsnicht.xlsx'))
    assert ergebnis['fehler']
    assert ergebnis['verschoben'] == []


def test_trockenlauf_schreibt_nicht(tmp_path):
    pfad = _datei(tmp_path, [
        [1, 'Wartung', 'P01', '02.11.2026', '02.11.2026', 'a', 'abgeschlossen'],
    ])
    ergebnis = verschiebe_abgeschlossene(pfad, trocken=True)
    assert ergebnis['verschoben'] == [1]
    wb = load_workbook(pfad)
    assert wb['Termine'].max_row == 2        # unveraendert
